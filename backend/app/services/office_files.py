"""Build real Office files from structured content the assistant supplies.

Word and PDF are sections of headings, paragraphs and bullets. Spreadsheets are named sheets of
columns and rows (or one sheet as CSV). Presentations are slides of titles and bullets, with
optional speaker notes. Charts, images, animations and spreadsheet formulas are not produced:
a cell that looks like a formula is stored as text so opening the file cannot run it.
"""

from __future__ import annotations

import csv
import io
import re
from pathlib import Path

from dhanz_knowledge.documents import safe_filename

_FONTS = Path(__file__).resolve().parent.parent / "assets" / "fonts"

MAX_SECTIONS = 40
MAX_PARAGRAPHS = 30
MAX_BULLETS = 40
MAX_SLIDES = 30
MAX_SHEETS = 8
MAX_COLUMNS = 40
MAX_ROWS = 500
MAX_TEXT = 4_000
_FORMULA = re.compile(r"^[=+@\t\r]")
_NUMBER = re.compile(r"-?(?:0|[1-9]\d*)(?:\.\d+)?$")
_SHEET_BAD = re.compile(r"[:\\/?*\[\]]")


class DocumentBuildError(Exception):
    """The requested document cannot be built; the message is safe to show."""


def filename_for(name: str, extension: str) -> str:
    stem = Path(safe_filename(name or "document")).stem.strip(" .") or "document"
    return f"{stem}.{extension}"[:150]


def build_docx(title: str, sections: list[dict]) -> bytes:
    import docx

    document = docx.Document()
    document.core_properties.creator = "DHANZ AI"
    document.core_properties.title = title
    document.add_heading(title, level=0)
    for section in sections:
        heading = section.get("heading") or ""
        if heading:
            document.add_heading(heading, level=1)
        for paragraph in section.get("paragraphs") or []:
            document.add_paragraph(paragraph)
        for bullet in section.get("bullets") or []:
            document.add_paragraph(bullet, style="List Bullet")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def build_pdf(title: str, sections: list[dict]) -> bytes:
    from fpdf import FPDF

    pdf = FPDF(format="A4")
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_font("DejaVu", "", str(_FONTS / "DejaVuSans.ttf"))
    pdf.add_font("DejaVu", "B", str(_FONTS / "DejaVuSans-Bold.ttf"))
    pdf.set_title(title)
    pdf.set_creator("DHANZ AI")
    pdf.add_page()
    width = pdf.epw
    pdf.set_font("DejaVu", "B", 16)
    pdf.multi_cell(width, 8, title, new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)
    for section in sections:
        heading = section.get("heading") or ""
        if heading:
            pdf.set_font("DejaVu", "B", 13)
            pdf.multi_cell(width, 7, heading, new_x="LMARGIN", new_y="NEXT")
        pdf.set_font("DejaVu", "", 11)
        for paragraph in section.get("paragraphs") or []:
            pdf.multi_cell(width, 6, paragraph, new_x="LMARGIN", new_y="NEXT")
            pdf.ln(1)
        for bullet in section.get("bullets") or []:
            pdf.multi_cell(width, 6, f"• {bullet}", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(2)
    return bytes(pdf.output())


def build_xlsx(sheets: list[dict]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.properties.creator = "DHANZ AI"
    used_titles: set[str] = set()
    for index, sheet in enumerate(sheets):
        worksheet = book.active if index == 0 else book.create_sheet()
        title = _sheet_title(sheet.get("name") or f"Sheet{index + 1}", index)
        worksheet.title = _unique_sheet_title(title, used_titles)
        columns = sheet["columns"]
        for column, name in enumerate(columns, start=1):
            cell = worksheet.cell(1, column, cell_value(name))
            cell.font = Font(bold=True)
        for row_index, row in enumerate(sheet["rows"], start=2):
            for column, value in enumerate(row, start=1):
                worksheet.cell(row_index, column, cell_value(value))
        worksheet.freeze_panes = "A2"
        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions
        for column in range(1, len(columns) + 1):
            widest = 0
            for row_index in range(1, worksheet.max_row + 1):
                widest = max(widest, len(str(worksheet.cell(row_index, column).value or "")))
            worksheet.column_dimensions[get_column_letter(column)].width = min(40, max(10, widest + 2))
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def build_csv(columns: list, rows: list[list]) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(cell_value(column) for column in columns)
    for row in rows:
        writer.writerow(cell_value(value) for value in row)
    return buffer.getvalue().encode("utf-8-sig")


def build_pptx(title: str, subtitle: str, slides: list[dict]) -> bytes:
    from pptx import Presentation

    presentation = Presentation()
    presentation.core_properties.creator = "DHANZ AI"
    presentation.core_properties.title = title
    opening = presentation.slides.add_slide(presentation.slide_layouts[0])
    opening.shapes.title.text = title
    if len(opening.placeholders) > 1:
        opening.placeholders[1].text = subtitle or "Created with DHANZ AI"
    for spec in slides:
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = spec["title"]
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = spec.get("bullets") or []
        if not bullets:
            frame.paragraphs[0].text = ""
        for index, bullet in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = bullet
            paragraph.level = 0
        notes = spec.get("notes") or ""
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def cell_value(value: object) -> str | int | float | bool:
    """A spreadsheet cell. Numbers stay numbers; anything that looks like a formula stays text."""
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):  # NaN or infinity
            raise DocumentBuildError("A spreadsheet cell is not a finite number.")
        return value
    text = "" if value is None else str(value).replace("\x00", "")
    text = text.strip()[:MAX_TEXT]
    if _NUMBER.fullmatch(text):
        return float(text) if "." in text else int(text)
    if text[:1] == "-" or _FORMULA.match(text):
        return "'" + text
    return text


def parse_sections(raw: object) -> list[dict]:
    if not isinstance(raw, list) or not raw:
        raise DocumentBuildError("sections must be a non-empty list.")
    if len(raw) > MAX_SECTIONS:
        raise DocumentBuildError(f"At most {MAX_SECTIONS} sections.")
    sections = []
    for item in raw:
        if not isinstance(item, dict):
            raise DocumentBuildError("Each section must be an object with heading, paragraphs and bullets.")
        heading = _text(item.get("heading"), "heading")
        paragraphs = _texts(item.get("paragraphs") or [], "paragraphs", MAX_PARAGRAPHS)
        bullets = _texts(item.get("bullets") or [], "bullets", MAX_BULLETS)
        if not heading and not paragraphs and not bullets:
            raise DocumentBuildError("Each section needs a heading, a paragraph or a bullet.")
        sections.append({"heading": heading, "paragraphs": paragraphs, "bullets": bullets})
    return sections


def parse_sheets(raw: object) -> list[dict]:
    if not isinstance(raw, list) or not raw:
        raise DocumentBuildError("sheets must be a non-empty list.")
    if len(raw) > MAX_SHEETS:
        raise DocumentBuildError(f"At most {MAX_SHEETS} sheets.")
    sheets = []
    for item in raw:
        if not isinstance(item, dict):
            raise DocumentBuildError("Each sheet must be an object with name, columns and rows.")
        columns = item.get("columns")
        rows = item.get("rows")
        if not isinstance(columns, list) or not columns:
            raise DocumentBuildError("Each sheet needs a columns list.")
        if len(columns) > MAX_COLUMNS:
            raise DocumentBuildError(f"At most {MAX_COLUMNS} columns.")
        if not isinstance(rows, list):
            raise DocumentBuildError("rows must be a list of rows.")
        if len(rows) > MAX_ROWS:
            raise DocumentBuildError(f"At most {MAX_ROWS} rows.")
        parsed_rows = []
        for row in rows:
            if not isinstance(row, list):
                raise DocumentBuildError("Each row must be a list of cell values, in column order.")
            if len(row) > MAX_COLUMNS:
                raise DocumentBuildError(f"At most {MAX_COLUMNS} columns.")
            parsed_rows.append(row)
        sheets.append({"name": _text(item.get("name"), "name") or "Sheet", "columns": columns, "rows": parsed_rows})
    return sheets


def parse_slides(raw: object) -> list[dict]:
    if not isinstance(raw, list) or not raw:
        raise DocumentBuildError("slides must be a non-empty list.")
    if len(raw) > MAX_SLIDES:
        raise DocumentBuildError(f"At most {MAX_SLIDES} slides.")
    slides = []
    for item in raw:
        if not isinstance(item, dict):
            raise DocumentBuildError("Each slide must be an object with title and bullets.")
        title = _text(item.get("title"), "title")
        if not title:
            raise DocumentBuildError("Each slide needs a title.")
        slides.append(
            {
                "title": title,
                "bullets": _texts(item.get("bullets") or [], "bullets", MAX_BULLETS),
                "notes": _text(item.get("notes"), "notes"),
            }
        )
    return slides


def _text(value: object, label: str) -> str:
    if value is None:
        return ""
    if not isinstance(value, str):
        raise DocumentBuildError(f"{label} must be text.")
    return value.replace("\x00", "").strip()[:MAX_TEXT]


def _texts(values: object, label: str, limit: int) -> list[str]:
    if not isinstance(values, list):
        raise DocumentBuildError(f"{label} must be a list of strings.")
    if len(values) > limit:
        raise DocumentBuildError(f"At most {limit} {label}.")
    texts = [_text(value, label) for value in values]
    return [text for text in texts if text]


def _sheet_title(name: str, index: int) -> str:
    cleaned = _SHEET_BAD.sub(" ", name).strip() or f"Sheet{index + 1}"
    return cleaned[:31]


def _unique_sheet_title(title: str, used: set[str]) -> str:
    candidate = title
    number = 2
    while candidate in used:
        suffix = f" {number}"
        candidate = f"{title[: 31 - len(suffix)]}{suffix}"
        number += 1
    used.add(candidate)
    return candidate
