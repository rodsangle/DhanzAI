"""Read a stored file back and describe it for an on-screen preview.

The preview is built from the bytes on disk, not from the text saved when the file was created,
so what the app shows is what a download would contain. Spreadsheets are cell grids, slides are
the text on each slide, Word files are headings, paragraphs, bullets and tables. PDF pages are
rendered to PNG. This is not a pixel-perfect Word or PowerPoint layout, and spreadsheet formulas
are shown as stored text rather than calculated.
"""

from __future__ import annotations

import csv
import io
from typing import Literal

MAX_ROWS = 200
MAX_COLUMNS = 40
MAX_BLOCKS = 400
MAX_SLIDES = 40
MAX_PDF_PAGES = 15
_PDF_SCALE = 1.5
_PDF_MAX_EDGE = 1400

PreviewKind = Literal["spreadsheet", "document", "slides", "pdf", "unsupported"]


class PreviewError(Exception):
    """The file cannot be previewed; the message is safe to show."""


def build_preview(filename: str, data: bytes) -> dict:
    ext = _extension(filename)
    if ext in {".xlsx", ".csv"}:
        return _spreadsheet(filename, data, ext)
    if ext == ".docx":
        return _document(filename, data)
    if ext == ".pptx":
        return _slides(filename, data)
    if ext == ".pdf":
        return _pdf_info(filename, data)
    return {
        "filename": filename,
        "file_type": ext.lstrip(".") or "file",
        "kind": "unsupported",
        "note": "This file type has no formatted preview.",
        "truncated": False,
        "sheets": [],
        "blocks": [],
        "slides": [],
        "page_count": 0,
    }


def render_pdf_page(data: bytes, page: int) -> bytes:
    """Render one 1-based PDF page to PNG."""
    if not data.startswith(b"%PDF-"):
        raise PreviewError("This file is not a PDF.")
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        count = len(pdf)
        if page < 1 or page > count:
            raise PreviewError("That page is not in the file.")
        if page > MAX_PDF_PAGES:
            raise PreviewError(f"Only the first {MAX_PDF_PAGES} pages can be previewed.")
        bitmap = pdf[page - 1].render(scale=_PDF_SCALE)
        image = bitmap.to_pil()
        image.thumbnail((_PDF_MAX_EDGE, _PDF_MAX_EDGE))
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        return buffer.getvalue()
    finally:
        pdf.close()


def _spreadsheet(filename: str, data: bytes, ext: str) -> dict:
    if ext == ".csv":
        sheets = [_csv_sheet(filename, data)]
        file_type = "csv"
    else:
        sheets = _xlsx_sheets(data)
        file_type = "xlsx"
    truncated = any(sheet.pop("_truncated") for sheet in sheets)
    note = f"Showing the first {MAX_ROWS} rows." if truncated else None
    return _envelope(filename, file_type, "spreadsheet", sheets=sheets, truncated=truncated, note=note)


def _xlsx_sheets(data: bytes) -> list[dict]:
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise PreviewError("This spreadsheet could not be previewed. It may be damaged.") from exc
    sheets = []
    try:
        for worksheet in book.worksheets:
            rows: list[list] = []
            truncated = False
            for index, row in enumerate(worksheet.iter_rows(values_only=True)):
                if index >= MAX_ROWS:
                    truncated = True
                    break
                values = [_cell(value) for value in row[:MAX_COLUMNS]]
                if any(value not in (None, "") for value in values):
                    rows.append(values)
            sheets.append({"name": worksheet.title, "rows": rows, "_truncated": truncated})
    finally:
        book.close()
    return sheets or [{"name": "Sheet", "rows": [], "_truncated": False}]


def _csv_sheet(filename: str, data: bytes) -> dict:
    text = data.decode("utf-8-sig", errors="replace")
    parsed = list(csv.reader(io.StringIO(text)))
    truncated = len(parsed) > MAX_ROWS
    rows = [[_cell(value) for value in row[:MAX_COLUMNS]] for row in parsed[:MAX_ROWS]]
    rows = [row for row in rows if any(value not in (None, "") for value in row)]
    name = filename.rsplit(".", 1)[0] or "Sheet"
    return {"name": name[:31], "rows": rows, "_truncated": truncated}


def _document(filename: str, data: bytes) -> dict:
    import docx
    from docx.opc.exceptions import PackageNotFoundError

    try:
        document = docx.Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, OSError) as exc:
        raise PreviewError("This document could not be previewed. It may be damaged or password-protected.") from exc
    blocks: list[dict] = []
    truncated = False
    for paragraph in document.paragraphs:
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        text = paragraph.text.strip()
        if not text:
            continue
        style = (paragraph.style.name or "").lower() if paragraph.style is not None else ""
        if style == "title" or style.startswith("heading"):
            digits = "".join(character for character in style if character.isdigit())
            blocks.append({"type": "heading", "level": int(digits) if digits else 1, "text": text})
        elif "list" in style:
            blocks.append({"type": "bullet", "text": text})
        else:
            blocks.append({"type": "paragraph", "text": text})
    for table in document.tables:
        if len(blocks) >= MAX_BLOCKS:
            truncated = True
            break
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if rows:
            blocks.append({"type": "table", "rows": rows[:MAX_ROWS]})
    note = "Showing the start of the document." if truncated else None
    return _envelope(filename, "docx", "document", blocks=blocks, truncated=truncated, note=note)


def _slides(filename: str, data: bytes) -> dict:
    from pptx import Presentation
    from pptx.exc import PackageNotFoundError

    try:
        presentation = Presentation(io.BytesIO(data))
    except (PackageNotFoundError, ValueError, OSError) as exc:
        raise PreviewError("This presentation could not be previewed. It may be damaged.") from exc
    slides = []
    truncated = len(presentation.slides) > MAX_SLIDES
    for number, slide in enumerate(presentation.slides, start=1):
        if number > MAX_SLIDES:
            break
        title = ""
        if slide.shapes.title is not None and slide.shapes.title.has_text_frame:
            title = slide.shapes.title.text_frame.text.strip()
        bullets: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            if slide.shapes.title is not None and shape == slide.shapes.title:
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text:
                    bullets.append(text)
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        slides.append({"number": number, "title": title or f"Slide {number}", "bullets": bullets, "notes": notes})
    note = f"Showing the first {MAX_SLIDES} slides." if truncated else None
    return _envelope(filename, "pptx", "slides", slides=slides, truncated=truncated, note=note)


def _pdf_info(filename: str, data: bytes) -> dict:
    if not data.startswith(b"%PDF-"):
        raise PreviewError("This file is not a PDF.")
    import pypdfium2 as pdfium

    pdf = pdfium.PdfDocument(data)
    try:
        count = len(pdf)
    finally:
        pdf.close()
    shown = min(count, MAX_PDF_PAGES)
    truncated = count > MAX_PDF_PAGES
    note = f"Showing the first {MAX_PDF_PAGES} pages." if truncated else None
    return _envelope(filename, "pdf", "pdf", page_count=shown, truncated=truncated, note=note)


def _envelope(
    filename: str,
    file_type: str,
    kind: PreviewKind,
    *,
    sheets: list[dict] | None = None,
    blocks: list[dict] | None = None,
    slides: list[dict] | None = None,
    page_count: int = 0,
    truncated: bool = False,
    note: str | None = None,
) -> dict:
    return {
        "filename": filename,
        "file_type": file_type,
        "kind": kind,
        "note": note,
        "truncated": truncated,
        "sheets": sheets or [],
        "blocks": blocks or [],
        "slides": slides or [],
        "page_count": page_count,
    }


def _extension(filename: str) -> str:
    dot = filename.rfind(".")
    return filename[dot:].lower() if dot >= 0 else ""


def _cell(value: object) -> str | int | float | bool | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return value
    if isinstance(value, int) and not isinstance(value, bool):
        return value
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        return value
    text = str(value).replace("\x00", "")
    return text
