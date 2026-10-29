"""Phase 3: file uploads & document intelligence, images, projects, memory, voice, PDF export."""

import base64
import io
from pathlib import Path

import httpx
import pytest
from fastapi.testclient import TestClient

from app.core.rate_limit import limiter
from app.main import create_app
from app.services.scanner import ClamAVScanner, MalwareFound, ScanError
from app.services.stt import SpeechToText
from dhanz_ai import Orchestrator
from tests.conftest import EchoProvider, parse_sse

PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def docx_bytes(text: str) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Maintenance manual", level=1)
    d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def xlsx_bytes() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(["Month", "Sales"])
    for month, value in [("Jan", 100), ("Feb", 250), ("Mar", 150)]:
        ws.append([month, value])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def upload(client, auth, name, data, **form):
    return client.post("/api/v1/files", headers=auth["headers"], files={"file": (name, data)}, data=form)


def chat(client, auth, content, conv_id=None, **extra):
    cid = conv_id or client.post("/api/v1/conversations", headers=auth["headers"], json={}).json()["id"]
    r = client.post(
        f"/api/v1/conversations/{cid}/messages", headers=auth["headers"], json={"content": content, **extra}
    )
    assert r.status_code == 200, r.text
    return cid, parse_sse(r.text)


def register(client, email):
    r = client.post("/api/v1/auth/register", json={"email": email, "password": "correct-horse-9", "display_name": "U"})
    return {"headers": {"Authorization": f"Bearer {r.json()['access_token']}"}}


# ---------------------------------------------------------------- files & documents


def test_upload_docx_download_and_library(client, auth, settings):
    data = docx_bytes("Replace the pump seal every 2,000 operating hours.")
    r = upload(client, auth, "manual.docx", data)
    assert r.status_code == 201, r.text
    doc = r.json()
    assert doc["kind"] == "file" and doc["file_type"] == "docx" and doc["scan_status"] == "not_scanned"
    assert doc["size_bytes"] == len(data) and doc["title"] == "manual.docx"

    d = client.get(f"/api/v1/files/{doc['id']}/download", headers=auth["headers"])
    assert d.content == data
    assert d.headers["content-type"] == "application/octet-stream"
    assert d.headers["content-disposition"].startswith("attachment;")

    listed = client.get("/api/v1/knowledge/documents", headers=auth["headers"], params={"kind": "file"}).json()
    assert listed["total"] == 1
    hits = client.get("/api/v1/knowledge/search", headers=auth["headers"], params={"q": "pump seal hours"}).json()
    assert hits[0]["document_id"] == doc["id"]
    assert len(list(Path(settings.storage_dir).rglob("*"))) >= 1


def test_attached_document_is_given_to_the_model(client, auth, provider):
    doc = upload(client, auth, "manual.docx", docx_bytes("Replace the pump seal every 2,000 operating hours.")).json()
    cid, events = chat(client, auth, "How often should the seal be replaced?", attachment_ids=[doc["id"]])
    assert events[0]["user_message"]["attachments"][0]["filename"] == "manual.docx"
    sources = next(e for e in events if e.get("stage") == "sources")["sources"]
    assert sources[0]["origin"] == "file" and sources[0]["url"] == f"upload:{doc['id']}"
    assert "every 2,000 operating hours" in provider.last_messages[0].content

    # Regenerate keeps using the attachment.
    r = client.post(f"/api/v1/conversations/{cid}/regenerate", headers=auth["headers"], json={})
    assert "every 2,000 operating hours" in provider.last_messages[0].content
    assert parse_sse(r.text)[-1]["type"] == "done"


def test_spreadsheet_statistics_reach_the_model(client, auth, provider):
    doc = upload(client, auth, "sales.xlsx", xlsx_bytes()).json()
    chat(client, auth, "What were total sales?", attachment_ids=[doc["id"]])
    assert "Sales: numeric, 3 values, sum=500" in provider.last_messages[0].content


def test_compare_two_documents(client, auth, provider):
    a = upload(client, auth, "a.txt", b"Offer A: price 1000 USD, delivery 4 weeks.").json()
    b = upload(client, auth, "b.txt", b"Offer B: price 900 USD, delivery 8 weeks.").json()
    chat(client, auth, "Compare these two offers", attachment_ids=[a["id"], b["id"]])
    system = provider.last_messages[0].content
    assert "[1] a.txt" in system and "[2] b.txt" in system


def test_upload_rejections(client, auth, settings):
    assert upload(client, auth, "virus.exe", b"MZ....").status_code == 422
    assert upload(client, auth, "fake.pdf", b"not a pdf").status_code == 422
    assert upload(client, auth, "empty.txt", b"").status_code == 422
    big = b"a" * (settings.max_upload_mb * 1024 * 1024 + 10)
    assert upload(client, auth, "big.txt", big).status_code == 413


def test_files_are_private(client, auth):
    doc = upload(client, auth, "secret.txt", b"The launch code is 1234.").json()
    eve = register(client, "eve@example.com")
    assert client.get(f"/api/v1/files/{doc['id']}/download", headers=eve["headers"]).status_code == 404
    cid = client.post("/api/v1/conversations", headers=eve["headers"], json={}).json()["id"]
    r = client.post(
        f"/api/v1/conversations/{cid}/messages",
        headers=eve["headers"],
        json={"content": "what does it say?", "attachment_ids": [doc["id"]]},
    )
    assert r.status_code == 404


def test_malware_scanner_verdicts(client, auth):
    class Fake:
        def __init__(self, error):
            self.error = error

        async def scan(self, data):
            if self.error:
                raise self.error

    files = client.app.state.files
    files.scanner = Fake(MalwareFound("Eicar-Test-Signature"))
    r = upload(client, auth, "eicar.txt", b"X5O!P%@AP")
    assert r.status_code == 422 and "Eicar-Test-Signature" in r.json()["detail"]
    files.scanner = Fake(ScanError("down"))
    assert upload(client, auth, "a.txt", b"hello there").status_code == 503
    files.scanner = Fake(None)
    assert upload(client, auth, "a.txt", b"hello there").json()["scan_status"] == "clean"


async def test_clamav_protocol():
    import asyncio

    async def serve(reply: bytes):
        async def handle(reader, writer):
            assert await reader.readexactly(10) == b"zINSTREAM\0"
            while True:
                size = int.from_bytes(await reader.readexactly(4), "big")
                if size == 0:
                    break
                await reader.readexactly(size)
            writer.write(reply)
            await writer.drain()
            writer.close()

        return await asyncio.start_server(handle, "127.0.0.1", 0)

    clean = await serve(b"stream: OK\0")
    await ClamAVScanner("127.0.0.1", clean.sockets[0].getsockname()[1]).scan(b"x" * 200_000)
    infected = await serve(b"stream: Eicar-Test-Signature FOUND\0")
    with pytest.raises(MalwareFound) as info:
        await ClamAVScanner("127.0.0.1", infected.sockets[0].getsockname()[1]).scan(b"x")
    assert info.value.signature == "Eicar-Test-Signature"
    with pytest.raises(ScanError):
        await ClamAVScanner("127.0.0.1", 1, timeout=1).scan(b"x")
    clean.close()
    infected.close()


# ---------------------------------------------------------------- images


def test_image_needs_a_vision_model(client, auth):
    img = upload(client, auth, "photo.png", PNG).json()
    assert img["file_type"] == "image"
    _, events = chat(client, auth, "What is in this picture?", attachment_ids=[img["id"]])
    assert events[-1]["type"] == "error" and "vision" in events[-1]["detail"]


def test_image_is_sent_to_vision_model(settings, knowledge):
    eye = EchoProvider("eye", vision=True)
    limiter.reset()
    with TestClient(create_app(settings, Orchestrator([EchoProvider(), eye]), knowledge)) as c:
        auth = register(c, "vision@example.com")
        features = c.get("/api/v1/models", headers=auth["headers"]).json()["features"]
        assert features["vision"] is True
        img = upload(c, auth, "photo.png", PNG).json()
        _, events = chat(c, auth, "What is in this picture?", attachment_ids=[img["id"]])
    assert events[-1]["type"] == "done" and events[-1]["message"]["model"] == "eye-1"
    last = eye.last_messages[-1]
    assert last.images[0].mime_type == "image/png" and base64.b64decode(last.images[0].data_b64) == PNG
    assert "[Attached images: photo.png]" in last.content


# ---------------------------------------------------------------- projects


def test_projects_group_chats_documents_and_instructions(client, auth, provider):
    h = auth["headers"]
    a = client.post("/api/v1/projects", headers=h, json={"name": "Plant A", "instructions": "Always use metric units."})
    b = client.post("/api/v1/projects", headers=h, json={"name": "Plant B"})
    pa, pb = a.json()["id"], b.json()["id"]

    upload(client, auth, "a-notes.txt", b"Plant A boiler pressure limit is 12 bar.", project_id=pa)
    conv_a = client.post("/api/v1/conversations", headers=h, json={"project_id": pa}).json()
    assert conv_a["project_id"] == pa
    chat(client, auth, "What is the boiler pressure limit?", conv_id=conv_a["id"])
    system = provider.last_messages[0].content
    assert "Always use metric units." in system and "12 bar" in system

    conv_b = client.post("/api/v1/conversations", headers=h, json={"project_id": pb}).json()
    chat(client, auth, "What is the boiler pressure limit?", conv_id=conv_b["id"])
    assert "12 bar" not in provider.last_messages[0].content  # project B can't see project A's documents

    listed = client.get("/api/v1/conversations", headers=h, params={"project_id": pa}).json()
    assert [c["id"] for c in listed["items"]] == [conv_a["id"]]
    detail = client.get(f"/api/v1/projects/{pa}", headers=h).json()
    assert detail["conversation_count"] == 1 and detail["document_count"] == 1

    note = upload(client, auth, "general.txt", b"General plant notes.").json()
    moved_doc = client.patch(f"/api/v1/knowledge/documents/{note['id']}", headers=h, json={"project_id": pb}).json()
    assert moved_doc["project_id"] == pb
    moved = client.patch(f"/api/v1/conversations/{conv_b['id']}", headers=h, json={"project_id": None}).json()
    assert moved["project_id"] is None
    assert client.patch(f"/api/v1/projects/{pa}", headers=h, json={"instructions": ""}).json()["instructions"] is None
    assert client.delete(f"/api/v1/projects/{pa}", headers=h).status_code == 204
    assert client.get(f"/api/v1/conversations/{conv_a['id']}", headers=h).json()["project_id"] is None

    eve = register(client, "eve@example.com")
    assert client.post("/api/v1/conversations", headers=eve["headers"], json={"project_id": pb}).status_code == 404
    assert client.get(f"/api/v1/projects/{pb}", headers=eve["headers"]).status_code == 404


# ---------------------------------------------------------------- memory


def test_memory_is_opt_in_and_used_when_enabled(client, auth, provider):
    h = auth["headers"]
    assert client.get("/api/v1/users/me", headers=h).json()["memory_enabled"] is False
    _, events = chat(client, auth, "Remember that I prefer metric units.")
    assert not any(e["type"] == "memory" for e in events)
    assert "memory is turned off" in provider.last_messages[0].content
    assert client.get("/api/v1/memories", headers=h).json() == []

    client.patch("/api/v1/users/me", headers=h, json={"memory_enabled": True})
    _, events = chat(client, auth, "Remember that I prefer metric units.")
    assert next(e for e in events if e["type"] == "memory")["content"] == "I prefer metric units"
    memories = client.get("/api/v1/memories", headers=h).json()
    assert [m["content"] for m in memories] == ["I prefer metric units"] and memories[0]["source"] == "chat"

    chat(client, auth, "How tall is Mount Everest?")
    assert "I prefer metric units" in provider.last_messages[0].content

    m = client.post("/api/v1/memories", headers=h, json={"content": "My name is Ada"}).json()
    assert (
        client.patch(f"/api/v1/memories/{m['id']}", headers=h, json={"content": "Call me Ada"}).json()["content"]
        == "Call me Ada"
    )
    assert len(client.get("/api/v1/users/me/export", headers=h).json()["memories"]) == 2

    client.patch("/api/v1/users/me", headers=h, json={"memory_enabled": False})
    chat(client, auth, "How tall is Mount Everest?")
    assert "metric units" not in provider.last_messages[0].content  # disabled = not used

    assert client.delete("/api/v1/memories", headers=h).status_code == 204
    assert client.get("/api/v1/memories", headers=h).json() == []
    eve = register(client, "eve@example.com")
    assert client.delete(f"/api/v1/memories/{m['id']}", headers=eve["headers"]).status_code == 404


# ---------------------------------------------------------------- voice


def test_transcription_not_configured(client, auth):
    r = client.post("/api/v1/voice/transcribe", headers=auth["headers"], files={"audio": ("a.m4a", b"x", "audio/m4a")})
    assert r.status_code == 501


def test_transcription(settings, provider, knowledge):
    seen = {}

    def handler(request: httpx.Request) -> httpx.Response:
        seen["url"] = str(request.url)
        seen["auth"] = request.headers.get("authorization")
        seen["body"] = request.content
        return httpx.Response(200, json={"text": " Hey Dhanz, what's the weather? "})

    stt = SpeechToText(
        "https://stt.example/v1", "sk-x", "whisper-1", httpx.AsyncClient(transport=httpx.MockTransport(handler))
    )
    limiter.reset()
    with TestClient(create_app(settings, Orchestrator([provider]), knowledge, stt)) as c:
        auth = register(c, "voice@example.com")
        assert c.get("/api/v1/models", headers=auth["headers"]).json()["features"]["speech_to_text"] is True
        r = c.post(
            "/api/v1/voice/transcribe",
            headers=auth["headers"],
            files={"audio": ("speech.m4a", b"\x00\x01audio", "audio/m4a")},
            data={"language": "en"},
        )
        assert r.json() == {"text": "Hey Dhanz, what's the weather?"}
        assert seen["url"] == "https://stt.example/v1/audio/transcriptions" and seen["auth"] == "Bearer sk-x"
        assert b"whisper-1" in seen["body"] and b"audio" in seen["body"]
        bad = c.post(
            "/api/v1/voice/transcribe", headers=auth["headers"], files={"audio": ("x.html", b"<b>", "text/html")}
        )
        assert bad.status_code == 415


# ---------------------------------------------------------------- export & deletion


def test_pdf_export(client, auth):
    cid, _ = chat(client, auth, "Привет! 你好 — explain ```python\nprint('hi')\n```")
    r = client.get(f"/api/v1/conversations/{cid}/export", headers=auth["headers"], params={"format": "pdf"})
    assert r.status_code == 200 and r.headers["content-type"] == "application/pdf"
    assert r.content.startswith(b"%PDF-") and len(r.content) > 1000


def test_account_deletion_removes_stored_files(client, auth, settings):
    upload(client, auth, "a.txt", b"some private notes here")
    stored = [p for p in Path(settings.storage_dir).rglob("*") if p.is_file()]
    assert stored
    r = client.request("DELETE", "/api/v1/users/me", headers=auth["headers"], json={"password": "correct-horse-9"})
    assert r.status_code == 204
    assert not any(p.exists() for p in stored)


def test_remember_request_parsing_and_redos_resistance():
    import time

    from app.services.chat_service import remember_request

    assert remember_request("Remember that I prefer metric units.") == "I prefer metric units"
    assert remember_request("  please REMEMBER: my sister is called Ana!  ") == "my sister is called Ana"
    assert remember_request("remember that\nI work nights") == "I work nights"
    assert remember_request("Do you remember that film?") is None
    assert remember_request("remember that ab") is None  # too short
    assert remember_request("remember that " + "x" * 600) is None  # too long for a memory

    # The old pattern took ~30 s on the first input (CodeQL py/polynomial-redos); messages may be 32,000 chars.
    for evil in (
        "remember that abc" + " " * 5_000 + "x",
        "remember that ab" + " " * 500 + "x",
        "remember:" + " " * 30_000,
    ):
        start = time.perf_counter()
        remember_request(evil)
        assert time.perf_counter() - start < 0.05
