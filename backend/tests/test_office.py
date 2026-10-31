"""Office files Dhanz creates: real Word, PDF, Excel, CSV and PowerPoint, saved to the library."""

import json
import uuid

import pytest

from app.services.assistant_tools import TOOLS, ToolBox, ToolContext
from app.services.office_files import (
    DocumentBuildError,
    build_csv,
    build_docx,
    build_pdf,
    build_pptx,
    build_xlsx,
    cell_value,
    parse_sections,
)
from dhanz_ai import ToolCall
from dhanz_knowledge.documents import extract_document


def test_spreadsheet_rows_declare_cell_types_for_strict_validators():
    rows = TOOLS["create_spreadsheet"].spec.parameters["properties"]["sheets"]["items"]["properties"]["rows"]
    assert rows["type"] == "array"
    assert rows["items"]["type"] == "array"
    assert rows["items"]["items"] == {"type": ["string", "number", "boolean", "null"]}


def _sections():
    return [
        {
            "heading": "Maintenance",
            "paragraphs": ["Replace the pump seal every 2,000 operating hours."],
            "bullets": ["Order the seal kit", "Record the hours"],
        }
    ]


def _sheets():
    return [{"name": "Budget", "columns": ["Item", "Amount"], "rows": [["Rent", 1200], ["Food", 80.5]]}]


async def _call(ctx: ToolContext, name: str, args: dict):
    box = ToolBox(ctx)
    result = await box.run(ToolCall("c1", name, args))
    return json.loads(result.content), result.is_error, box.activity


def _call_on_app_loop(client, ctx: ToolContext, name: str, args: dict):
    """Run the tool on the API's event loop.

    TestClient owns that loop. asyncpg connections are bound to it, so awaiting the
    tool from pytest's loop fails on PostgreSQL with "attached to a different loop".
    """
    return client.portal.call(_call, ctx, name, args)


def test_builders_round_trip_through_the_same_reader_as_uploads():
    word = extract_document("notes.docx", build_docx("Plant notes", _sections()))
    assert "Maintenance" in word.text and "pump seal" in word.text and "Order the seal kit" in word.text

    pdf = extract_document("notes.pdf", build_pdf("Plant notes", _sections()))
    assert pdf.kind == "pdf" and "pump seal" in pdf.text and "Order the seal kit" in pdf.text

    book = extract_document("budget.xlsx", build_xlsx(_sheets()))
    assert "Rent" in book.text and "1200" in book.text and "80.5" in book.text

    table = extract_document("budget.csv", build_csv(_sheets()[0]["columns"], _sheets()[0]["rows"]))
    assert "Food" in table.text and "80.5" in table.text

    slides = extract_document(
        "pitch.pptx",
        build_pptx(
            "DHANZ", "A personal assistant", [{"title": "Problem", "bullets": ["Too many tabs"], "notes": "Pause"}]
        ),
    )
    assert "Problem" in slides.text and "Too many tabs" in slides.text and "Speaker notes: Pause" in slides.text


def test_spreadsheet_formulas_are_stored_as_text():
    from io import BytesIO

    from openpyxl import load_workbook

    sheet = build_xlsx([{"name": "Q", "columns": ["Note"], "rows": [["=cmd|' /C calc'!A0"], ["+1+1"], ["@SUM(A1)"]]}])
    book = load_workbook(BytesIO(sheet))
    values = [row[0] for row in book.active.iter_rows(min_row=2, values_only=True)]
    # A leading apostrophe makes Excel keep the cell as text instead of running a formula.
    assert values == ["'=cmd|' /C calc'!A0", "'+1+1", "'@SUM(A1)"]


def test_cell_values_keep_numbers_and_reject_non_finite():
    assert cell_value(1200) == 1200
    assert cell_value("80.5") == 80.5
    assert cell_value("-3") == -3
    assert cell_value(True) is True
    with pytest.raises(DocumentBuildError):
        cell_value(float("nan"))


def test_sections_must_contain_something():
    with pytest.raises(DocumentBuildError):
        parse_sections([{"heading": "", "paragraphs": [], "bullets": []}])


def test_documents_are_saved_and_only_the_owner_can_download(client, auth):
    app = client.app
    other = client.post(
        "/api/v1/auth/register",
        json={"email": "other@example.com", "password": "correct-horse-9", "display_name": "Other"},
    )
    assert other.status_code == 201
    ctx = ToolContext(
        app.state.sessionmaker,
        uuid.UUID(auth["user"]["id"]),
        None,
        files=app.state.files,
        embedder=app.state.knowledge.embedder,
        settings=app.state.settings,
    )
    body, failed, activity = _call_on_app_loop(
        client,
        ctx,
        "create_word_document",
        {"title": "Plant notes", "filename": "Plant notes", "sections": _sections()},
    )
    assert failed is False and body["saved"] is True and body["filename"] == "Plant notes.docx"
    assert activity[0]["action"]["type"] == "file"

    download = client.get(f"/api/v1/files/{body['document_id']}/download", headers=auth["headers"])
    assert download.status_code == 200 and download.content.startswith(b"PK")
    assert "pump seal" in extract_document("Plant notes.docx", download.content).text

    library = client.get("/api/v1/knowledge/documents", headers=auth["headers"])
    assert library.status_code == 200
    match = next(item for item in library.json()["items"] if item["id"] == body["document_id"])
    assert match["added_via"] == "generated" and match["file_type"] == "docx"

    stolen = client.get(
        f"/api/v1/files/{body['document_id']}/download",
        headers={"Authorization": f"Bearer {other.json()['access_token']}"},
    )
    assert stolen.status_code == 404

    sheet, sheet_failed, _ = _call_on_app_loop(
        client, ctx, "create_spreadsheet", {"filename": "Budget", "sheets": _sheets()}
    )
    assert sheet_failed is False and sheet["file_type"] == "xlsx"
    csv_file, csv_failed, _ = _call_on_app_loop(
        client, ctx, "create_spreadsheet", {"filename": "Budget", "format": "csv", "sheets": _sheets()}
    )
    assert csv_failed is False and csv_file["filename"].endswith(".csv")
    deck, deck_failed, _ = _call_on_app_loop(
        client,
        ctx,
        "create_presentation",
        {"title": "Pitch", "slides": [{"title": "Problem", "bullets": ["Too many tabs"], "notes": "Pause"}]},
    )
    assert deck_failed is False and deck["file_type"] == "pptx"
    pdf, pdf_failed, _ = _call_on_app_loop(
        client, ctx, "create_word_document", {"title": "Plant notes", "format": "pdf", "sections": _sections()}
    )
    assert pdf_failed is False and pdf["filename"].endswith(".pdf")
    pdf_bytes = client.get(f"/api/v1/files/{pdf['document_id']}/download", headers=auth["headers"]).content
    assert pdf_bytes.startswith(b"%PDF-") and "pump seal" in extract_document("notes.pdf", pdf_bytes).text

    def preview(document_id: str) -> dict:
        response = client.get(f"/api/v1/files/{document_id}/preview", headers=auth["headers"])
        assert response.status_code == 200, response.text
        return response.json()

    word_view = preview(body["document_id"])
    assert word_view["kind"] == "document"
    assert any(block["type"] == "heading" and block["text"] == "Maintenance" for block in word_view["blocks"])
    assert any("pump seal" in block["text"] for block in word_view["blocks"])
    assert any(block["type"] == "bullet" and block["text"] == "Order the seal kit" for block in word_view["blocks"])

    sheet_view = preview(sheet["document_id"])
    assert sheet_view["kind"] == "spreadsheet"
    assert sheet_view["sheets"][0]["rows"][0] == ["Item", "Amount"]
    assert sheet_view["sheets"][0]["rows"][1] == ["Rent", 1200]

    csv_view = preview(csv_file["document_id"])
    assert csv_view["file_type"] == "csv" and csv_view["sheets"][0]["rows"][2][0] == "Food"

    deck_view = preview(deck["document_id"])
    titles = [slide["title"] for slide in deck_view["slides"]]
    assert titles[:2] == ["Pitch", "Problem"]
    assert "Too many tabs" in deck_view["slides"][1]["bullets"]
    assert deck_view["slides"][1]["notes"] == "Pause"

    pdf_view = preview(pdf["document_id"])
    assert pdf_view["kind"] == "pdf" and pdf_view["page_count"] >= 1
    page = client.get(f"/api/v1/files/{pdf['document_id']}/preview/pages/1", headers=auth["headers"])
    assert page.status_code == 200 and page.headers["content-type"].startswith("image/png")
    assert page.content.startswith(b"\x89PNG")
    missing = client.get(f"/api/v1/files/{pdf['document_id']}/preview/pages/99", headers=auth["headers"])
    assert missing.status_code == 404
    not_pdf = client.get(f"/api/v1/files/{body['document_id']}/preview/pages/1", headers=auth["headers"])
    assert not_pdf.status_code == 422

    other_headers = {"Authorization": f"Bearer {other.json()['access_token']}"}
    assert client.get(f"/api/v1/files/{body['document_id']}/preview", headers=other_headers).status_code == 404
    assert client.get(f"/api/v1/files/{pdf['document_id']}/preview/pages/1", headers=other_headers).status_code == 404


async def test_document_tools_report_bad_input_and_missing_storage():
    app_ctx = ToolContext(None, uuid.uuid4(), None)
    body, failed, activity = await _call(
        app_ctx, "create_spreadsheet", {"filename": "Budget", "format": "csv", "sheets": _sheets() * 2}
    )
    assert failed is True and "one sheet" in body["error"]
    assert activity[0]["error"] is True

    bare = ToolContext(None, uuid.uuid4(), None)
    body, failed, _ = await _call(
        bare, "create_word_document", {"title": "Notes", "sections": [{"paragraphs": ["Hello"]}]}
    )
    assert failed is True and "not available" in body["error"]
