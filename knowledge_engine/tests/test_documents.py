import io
import zipfile

import pytest

from dhanz_knowledge.documents import UnsupportedFile, describe_table, extract_document, safe_filename


def make_docx() -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("Pump specification", level=1)
    d.add_paragraph("The pump delivers 40 m3/h at 25 m head.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Item", "Value"
    t.cell(1, 0).text, t.cell(1, 1).text = "Power", "5.5 kW"
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def make_xlsx() -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sales"
    ws.append(["Region", "Revenue"])
    for region, rev in [("North", 1200), ("South", 800), ("East", 1000)]:
        ws.append([region, rev])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def make_pptx() -> bytes:
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    slide.placeholders[1].text = "Revenue grew 12%"
    slide.notes_slide.notes_text_frame.text = "Mention the new plant"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def make_pdf(text: str | None = "Invoice total: 1,250 USD") -> bytes:
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    if text:
        pdf.set_font("Helvetica", size=12)
        pdf.cell(0, 10, text)
    return bytes(pdf.output())


def test_docx_headings_paragraphs_and_tables():
    doc = extract_document("spec.docx", make_docx())
    assert doc.kind == "docx"
    assert "## Pump specification" in doc.text and "40 m3/h" in doc.text
    assert "| Power | 5.5 kW |" in doc.text


def test_xlsx_statistics_are_exact():
    doc = extract_document("sales.xlsx", make_xlsx())
    assert "## Sheet: Sales" in doc.text and "3 data rows × 2 columns" in doc.text
    assert "Revenue: numeric, 3 values, sum=3,000, mean=1,000, min=800, max=1,200" in doc.text
    assert "| North | 1200 |" in doc.text


def test_csv_with_semicolons():
    doc = extract_document("data.csv", b"name;score\nada;9\nbob;7\n")
    assert doc.kind == "csv" and "score: numeric, 2 values, sum=16" in doc.text


def test_pptx_slides_and_notes():
    doc = extract_document("deck.pptx", make_pptx())
    assert (
        "## Slide 1" in doc.text
        and "Revenue grew 12%" in doc.text
        and "Speaker notes: Mention the new plant" in doc.text
    )


def test_pdf_text_and_scanned_pdf_note():
    assert "Invoice total: 1,250 USD" in extract_document("inv.pdf", make_pdf()).text
    scanned = extract_document("scan.pdf", make_pdf(None))
    assert scanned.text == "" and "OCR" in scanned.notes[0]


def test_text_code_and_images():
    assert extract_document("main.py", b"print('hi')\n").kind == "code"
    assert extract_document("notes.md", "# Café".encode()).text == "# Café"
    png = b"\x89PNG\r\n\x1a\n" + b"\x00" * 20
    img = extract_document("photo.PNG", png)
    assert img.is_image and img.mime_type == "image/png"


@pytest.mark.parametrize(
    "name,data,message",
    [
        ("evil.pdf", b"MZ\x90\x00 not a pdf", "valid PDF"),
        ("fake.png", b"<svg onload=alert(1)>", "not really an image"),
        ("fake.docx", b"%PDF-1.4", "valid Office"),
        ("old.doc", b"\xd0\xcf\x11\xe0", "Legacy .doc"),
        ("app.exe", b"MZ", "not supported"),
        ("bin.txt", b"\x00\x01\x02binary", "binary"),
    ],
)
def test_rejections(name, data, message):
    with pytest.raises(UnsupportedFile, match=message):
        extract_document(name, data)


def test_zip_bomb_rejected():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"0" * 120_000_000)
    with pytest.raises(UnsupportedFile, match="unreasonable size"):
        extract_document("bomb.docx", buf.getvalue())


def test_safe_filename_and_table_text_columns():
    assert safe_filename("../../etc/passwd") == "passwd"
    assert safe_filename("C:\\Users\\x\\My Report?.pdf") == "My Report_.pdf"
    out = describe_table([["city", "n"], ["Paris", "x"], ["Rome", "y"]])
    assert "city: text, 2 values, 2 distinct" in out
